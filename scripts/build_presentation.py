from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "reports" / "figures"
OUTPUT = ROOT / "reports" / "presentation.pptx"

NAVY = RGBColor(18, 42, 62)
TEAL = RGBColor(20, 126, 139)
INK = RGBColor(41, 57, 69)
MUTED = RGBColor(99, 116, 128)
PALE = RGBColor(241, 246, 247)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(223, 157, 58)
BLUE_PALE = RGBColor(226, 239, 242)
FONT = "Aptos"
SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5


def arabic(text: str) -> str:
    return get_display(arabic_reshaper.reshape(text))


def add_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=INK,
    bold=False,
    rtl=True,
    align=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align or (PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)
    run = paragraph.add_run()
    run.text = arabic(text) if rtl else text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, fill=PALE, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
    return shape


def add_header(slide, section, title, subtitle=None, page=None, source=None):
    add_text(slide, section.upper(), 0.72, 0.24, 11.9, 0.28, 10, TEAL, True)
    add_text(slide, title, 0.72, 0.62, 11.9, 0.52, 25, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, 0.75, 1.15, 11.8, 0.42, 12, MUTED)
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.62), Inches(0.72), Inches(0.045)
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = TEAL
    footer = f"تحليل الطلب على الوجبات  |  {page:02d}" if page else "تحليل الطلب على الوجبات"
    if source:
        footer += f"  •  {source}"
    add_text(slide, footer, 0.72, 7.12, 11.95, 0.2, 8, MUTED)


def add_metric(slide, x, y, w, value, label, detail=None, accent=TEAL):
    add_card(slide, x, y, w, 1.25, WHITE, line=BLUE_PALE)
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + w - 0.08), Inches(y + 0.16), Inches(0.035), Inches(0.92)
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = accent
    add_text(slide, value, x + 0.16, y + 0.13, w - 0.38, 0.48, 25, NAVY, True)
    add_text(slide, label, x + 0.16, y + 0.62, w - 0.38, 0.32, 12, INK, True)
    if detail:
        add_text(slide, detail, x + 0.16, y + 0.93, w - 0.38, 0.22, 9, MUTED)


def add_bullet(slide, text, x, y, w, h=0.65, size=15, color=INK, marker=TEAL):
    dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + w - 0.12), Inches(y + 0.18), Inches(0.075), Inches(0.075)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = marker
    dot.line.fill.background()
    return add_text(slide, text, x, y, w - 0.2, h, size, color)


def add_figure(slide, filename, x, y, w, h):
    path = FIGURES / filename
    if not path.is_file():
        raise FileNotFoundError(f"Run the analysis notebook first; missing {path}")
    with Image.open(path) as image:
        image_w, image_h = image.size
    scale = min(w / image_w, h / image_h)
    draw_w, draw_h = image_w * scale, image_h * scale
    left = x + (w - draw_w) / 2
    top = y + (h - draw_h) / 2
    slide.shapes.add_picture(
        str(path), Inches(left), Inches(top), width=Inches(draw_w), height=Inches(draw_h)
    )


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(deck, section, title, subtitle, page, source):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_background(slide)
    add_header(slide, section, title, subtitle, page, source)
    return slide


def main():
    if not FIGURES.is_dir():
        raise FileNotFoundError("Run notebooks/analysis.ipynb to generate the charts first.")

    deck = Presentation()
    deck.slide_width = Inches(SLIDE_WIDTH)
    deck.slide_height = Inches(SLIDE_HEIGHT)
    blank = deck.slide_layouts[6]

    slide = deck.slides.add_slide(blank)
    add_background(slide, NAVY)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(SLIDE_HEIGHT)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_text(slide, "مشروع تحليل بيانات الطلب على الوجبات", 0.95, 0.9, 11.4, 0.55, 15, GOLD, True)
    add_text(slide, "كيف نخطط للطلب بذكاء؟", 0.95, 1.55, 11.2, 0.8, 34, WHITE, True)
    add_text(
        slide,
        "تحليل بيانات أسبوعية لشركة توصيل وجبات لفهم تغير الطلب واختلافه بين الوجبات ومراكز التجهيز، ودراسة ارتباطه بالعروض والخصومات لدعم تخطيط المخزون.",
        1.0,
        2.55,
        10.9,
        0.95,
        16,
        RGBColor(195, 216, 222),
    )
    add_card(slide, 0.98, 4.0, 11.15, 1.2, RGBColor(26, 57, 76))
    add_text(slide, "119.56 مليون", 1.25, 4.12, 3.25, 0.48, 25, WHITE, True)
    add_text(slide, "طلب مسجل", 1.3, 4.65, 3.0, 0.28, 11, RGBColor(195, 216, 222))
    add_text(slide, "145 أسبوعًا", 5.0, 4.12, 2.8, 0.48, 23, WHITE, True)
    add_text(slide, "في بيانات التدريب", 5.05, 4.65, 2.7, 0.28, 11, RGBColor(195, 216, 222))
    add_text(slide, "77 مركزًا", 8.5, 4.12, 2.8, 0.48, 23, WHITE, True)
    add_text(slide, "تجهيز", 8.55, 4.65, 2.7, 0.28, 11, RGBColor(195, 216, 222))
    add_text(slide, "إعداد: روان مسلم", 1.0, 6.2, 10.8, 0.42, 15, WHITE, True)
    add_text(slide, "تحليل وصفي لدعم قرارات أفضل — الارتباط لا يثبت السببية", 1.0, 6.7, 10.8, 0.32, 10, RGBColor(195, 216, 222))
    add_notes(
        slide,
        "افتتاحية: هذا المشروع يحلل بيانات أسبوعية مجهولة الهوية لشركة توصيل وجبات. ندرس كيف يتغير الطلب بين الأسابيع والفئات والوجبات ومراكز التجهيز، "
        "ونفحص ارتباط الطلب بالعروض والخصومات، بهدف دعم تخطيط المخزون والتسويق. "
        "الهدف هو تحويل النتائج إلى خطوات قابلة للاختبار، لا الادعاء بأن البيانات تثبت سبب تغير الطلب. الأرقام مأخوذة من خلايا الدفتر 2 و6.",
    )

    slide = new_slide(
        deck,
        "01 | الإطار",
        "ما السؤال الذي نريد أن نجيب عنه؟",
        "نحوّل مشكلة التخطيط والتسويق إلى أسئلة قابلة للقياس في بيانات الطلب.",
        2,
        "خلايا الدفتر 1–2",
    )
    add_card(slide, 0.8, 1.95, 5.75, 4.55, PALE)
    add_text(slide, "تحدي العمل", 1.1, 2.2, 4.95, 0.4, 18, NAVY, True)
    add_bullet(slide, "الطلب يتغير بين الأسابيع والوجبات والمراكز.", 1.05, 2.85, 5.0)
    add_bullet(slide, "العروض والأسعار تتغير، لكن بياناتنا رصدية.", 1.05, 3.7, 5.0)
    add_bullet(slide, "غياب سجل قد يعني عدم توفر البيانات، لا طلبًا صفريًا.", 1.05, 4.55, 5.0)
    add_bullet(slide, "لا تتوفر بيانات المخزون والهدر والتكلفة.", 1.05, 5.4, 5.0)
    add_card(slide, 6.8, 1.95, 5.72, 4.55, WHITE, line=BLUE_PALE)
    add_text(slide, "مسار التحليل", 7.1, 2.2, 4.9, 0.4, 18, NAVY, True)
    steps = [
        "١  فحص الملفات والمفاتيح وجودة البيانات",
        "٢  ربط الأبعاد وحساب مؤشرات الطلب",
        "٣  استكشاف الاتجاه والفئات والمراكز",
        "٤  اختبار الارتباط والتذبذب والتنبؤ",
        "٥  تحويل النتائج إلى تجربة قابلة للقياس",
    ]
    for i, item in enumerate(steps):
        add_bullet(slide, item, 7.05, 2.8 + i * 0.65, 5.0, h=0.52, size=13, marker=GOLD if i == 4 else TEAL)
    add_notes(
        slide,
        "قدّم المشكلة قبل الأرقام: وصف الشركة في المشروع يذكر الهدر والنفاد، لكن ملفاتنا لا تتضمن قياسات مباشرة لهما. "
        "لذلك نحلل الطلب المسجل كمؤشر تشغيلي محدود. اشرح ترتيب الدفتر: تحميل البيانات والتحقق، التنظيف، الربط، المؤشرات، ثم أسئلة العمل الستة. "
        "الخلايا المرجعية: 1–2.",
    )

    slide = new_slide(
        deck,
        "02 | البيانات والجودة",
        "قبل التحليل: هل البيانات صالحة للربط والمقارنة؟",
        "ثبّتنا المفاتيح، راجعنا القيم غير المعتادة، وحافظنا على السجلات الأصلية.",
        3,
        "خلايا الدفتر 2–5",
    )
    add_metric(slide, 0.85, 1.95, 2.75, "456,548", "سجل تدريب")
    add_metric(slide, 3.8, 1.95, 2.75, "32,573", "سجل اختبار", "لا يحتوي هدف الطلب")
    add_metric(slide, 6.75, 1.95, 2.75, "77 + 51", "مركزًا + وجبة")
    add_metric(slide, 9.7, 1.95, 2.75, "0", "مفاتيح تدريب مكررة")
    add_card(slide, 0.85, 3.55, 5.75, 2.75, PALE)
    add_text(slide, "فحوص وقرارات", 1.15, 3.8, 4.9, 0.35, 16, NAVY, True)
    add_bullet(slide, "لا قيم مفقودة أو أسابيع عامة مفقودة.", 1.1, 4.35, 5.0, h=0.45, size=12)
    add_bullet(slide, "2,350 سلسلة مركز–وجبة غير مكتملة؛ لا نملأها بأصفار.", 1.1, 4.9, 5.0, h=0.55, size=12)
    add_bullet(slide, "116,101 سعر بيع أعلى من الأساسي؛ أُبقيت دون افتراض أنها أخطاء.", 1.1, 5.55, 5.0, h=0.55, size=11)
    add_card(slide, 6.8, 3.55, 5.65, 2.75, WHITE, line=BLUE_PALE)
    add_text(slide, "سلامة الدمج", 7.1, 3.8, 4.85, 0.35, 16, NAVY, True)
    add_text(slide, "ربط many-to-one بمفاتيح الأبعاد الفريدة", 7.1, 4.35, 4.85, 0.48, 13, INK, True)
    add_text(slide, "456,548 صفًا قبل الربط وبعده", 7.1, 4.95, 4.85, 0.42, 13, TEAL, True)
    add_text(slide, "119,557,485 طلبًا قبل الربط وبعده", 7.1, 5.5, 4.85, 0.42, 13, TEAL, True)
    add_notes(
        slide,
        "اشرح أن train يحتوي الهدف num_orders، بينما test لا يحتويه ويغطي أسابيع لاحقة. تم فحص فرادة مفاتيح المركز والوجبة ومفتاح الأسبوع/المركز/الوجبة. "
        "بعد دمج جداول الأبعاد، بقي عدد الصفوف وإجمالي الطلب ثابتين. لا تعني السلسلة غير المكتملة أن الطلب صفر، ولم تُحذف القيم الشاذة تلقائيًا. "
        "سجل كل قرار وسببه في docs/cleaning_decisions.csv. الخلايا 2–5.",
    )

    slide = new_slide(
        deck,
        "03 | مؤشرات الأداء",
        "ملخص الطلب في لوحة واحدة",
        "المؤشرات محسوبة على سجلات التدريب؛ الإيراد تقديري وليس ربحًا.",
        4,
        "خلية الدفتر 6",
    )
    add_metric(slide, 0.85, 2.0, 2.75, "119.56M", "إجمالي الطلبات", "145 أسبوعًا")
    add_metric(slide, 3.8, 2.0, 2.75, "824,534", "متوسط الطلب الأسبوعي")
    add_metric(slide, 6.75, 2.0, 2.75, "33.9%", "حصة المشروبات", "من الطلب المسجل")
    add_metric(slide, 9.7, 2.0, 2.75, "22 من 51", "وجبة لبلوغ 80%", "من الطلب تراكميًا")
    add_card(slide, 0.85, 3.8, 11.6, 1.95, PALE)
    add_text(slide, "ماذا يدخل في المؤشرات؟", 1.15, 4.05, 10.7, 0.36, 16, NAVY, True)
    add_bullet(slide, "الإيراد التقديري = سعر البيع × الطلبات؛ لا يشمل تكاليف المواد أو التشغيل.", 1.1, 4.55, 10.8, h=0.46, size=13)
    add_bullet(slide, "الخصم = (السعر الأساسي − سعر البيع) ÷ السعر الأساسي؛ قد تكون النسبة سالبة.", 1.1, 5.05, 10.8, h=0.46, size=13)
    add_bullet(slide, "مؤشر الترويج يسجل وجود البريد أو الظهور على الصفحة، لا أثرهما السببي.", 1.1, 5.55, 10.8, h=0.46, size=13)
    add_notes(
        slide,
        "عرّف المقاييس كي لا تختلط على الجمهور: إجمالي الطلبات هو مجموع num_orders، ومتوسط الأسبوع هو متوسط المجاميع الأسبوعية. "
        "الإيراد التقديري حاصل ضرب checkout_price في الطلب، ولا يمثل الربح. نسبة الخصم مشتقة من base_price وcheckout_price؛ قد تصبح سالبة عندما يتجاوز سعر البيع السعر الأساسي. "
        "نسبة صفوف الترويج مؤشر وصفي. النتائج من خلية 6، وحصة الفئة وعدد الوجبات من خلية 8.",
    )

    slide = new_slide(
        deck,
        "04 | السؤال 1 — الاتجاه الأسبوعي",
        "الطلب يتقلب بوضوح بين الأسابيع",
        "الخط يمثل مجموع الطلبات المسجلة في كل أسبوع؛ أرقام الأسابيع ليست تواريخ تقويمية.",
        5,
        "خلية الدفتر 7",
    )
    add_card(slide, 0.8, 1.9, 8.05, 4.75, WHITE, line=BLUE_PALE)
    add_figure(slide, "01_weekly_demand.png", 1.05, 2.15, 7.55, 4.15)
    add_metric(slide, 9.15, 2.1, 3.15, "1.30M", "أعلى أسبوع", "الأسبوع 48", accent=GOLD)
    add_metric(slide, 9.15, 3.7, 3.15, "380,065", "أدنى أسبوع", "الأسبوع 62")
    add_card(slide, 9.15, 5.3, 3.15, 1.15, PALE)
    add_text(slide, "الأعلى ≈ 243% فوق الأدنى", 9.35, 5.52, 2.75, 0.65, 13, NAVY, True)
    add_notes(
        slide,
        "ابدأ بقراءة المحورين: كل نقطة إجمالي الطلب المسجل في أسبوع. أعلى قيمة 1,303,457 في الأسبوع 48، وأدناها 380,065 في الأسبوع 62؛ "
        "الأعلى أعلى بنحو 243% من الأدنى. هذه مقارنة داخل العينة وليست توقعًا أو دليلًا على الموسمية؛ لا توجد تواريخ فعلية. "
        "الخلايا 7 وملخص النتائج في 14.",
    )

    slide = new_slide(
        deck,
        "05 | السؤال 2 — أين يتركز الطلب؟",
        "المشروبات تقود الفئات؛ والطلب يتركز في مجموعة وجبات",
        "الترتيب مبني على مجموع الطلب، لا على الربحية أو الهامش.",
        6,
        "خلية الدفتر 8",
    )
    add_card(slide, 0.8, 1.9, 7.45, 4.75, WHITE, line=BLUE_PALE)
    add_figure(slide, "02_demand_by_category.png", 1.05, 2.15, 6.95, 4.15)
    add_metric(slide, 8.7, 2.0, 3.55, "33.9%", "حصة فئة المشروبات", "من كل الطلب المسجل")
    add_metric(slide, 8.7, 3.55, 3.55, "9.77M", "الوجبة الأعلى: 2290", "طلبات تراكمية")
    add_card(slide, 8.7, 5.1, 3.55, 1.3, PALE)
    add_text(slide, "22 وجبة من 51 تمثل 80% من الطلب", 8.95, 5.35, 3.05, 0.72, 14, NAVY, True)
    add_notes(
        slide,
        "اشرح طريقتين للتجميع: حسب category لمعرفة حصة الفئة، وحسب meal_id لترتيب الوجبات وحساب الحصة التراكمية. "
        "Beverages تمثل 33.9%، والوجبة 2290 أعلى مجموعًا بنحو 9.77 مليون طلب. احتجنا إلى 22 وجبة من أصل 51 لتجاوز 80%. "
        "هذه حصص طلب مسجل وليست أرباحًا. الخلية 8.",
    )

    slide = new_slide(
        deck,
        "06 | السؤال 3 — المراكز",
        "إجمالي الطلب وحده لا يعني أن المركز أكفأ",
        "نوع المركز وحجمه وتغطيته يؤثرون في قراءة المجاميع.",
        7,
        "خلية الدفتر 9",
    )
    add_card(slide, 0.8, 1.9, 7.45, 4.75, WHITE, line=BLUE_PALE)
    add_figure(slide, "03_center_type.png", 1.05, 2.15, 6.95, 4.15)
    add_metric(slide, 8.7, 2.0, 3.55, "4.30M", "أعلى مركز: 13", "TYPE_B")
    add_card(slide, 8.7, 3.65, 3.55, 2.55, PALE)
    add_text(slide, "متوسط الطلب لكل سجل", 8.95, 3.9, 3.0, 0.38, 14, NAVY, True)
    add_text(slide, "TYPE_A     262", 8.95, 4.5, 3.0, 0.34, 13, INK)
    add_text(slide, "TYPE_B     319", 8.95, 4.98, 3.0, 0.34, 13, INK)
    add_text(slide, "TYPE_C     207", 8.95, 5.46, 3.0, 0.34, 13, INK)
    add_notes(
        slide,
        "الرسم يقارن مجموع الطلب حسب نوع المركز، والجدول يعرض عدد المراكز ومتوسط الطلب لكل سجل. المركز 13 من النوع TYPE_B سجّل أعلى مجموع، نحو 4.30 مليون. "
        "متوسطات الطلب المسجل لكل سجل تقريبًا 262 في TYPE_A و319 في TYPE_B و207 في TYPE_C. لا نسمّي ذلك كفاءة؛ فالأحجام وعدد المراكز والتغطية تختلف. "
        "الخلايا 9 و14.",
    )

    slide = new_slide(
        deck,
        "07 | السؤال 4 — الترويج",
        "سجلات الترويج ارتبطت بوسيط طلب أعلى",
        "مقارنة داخل أزواج المركز–الوجبة التي ظهرت في حالتي الترويج وعدم الترويج.",
        8,
        "خلية الدفتر 10",
    )
    add_card(slide, 0.8, 1.9, 7.25, 4.75, WHITE, line=BLUE_PALE)
    add_figure(slide, "04_promotion_comparison.png", 1.05, 2.15, 6.75, 4.15)
    add_metric(slide, 8.5, 2.0, 3.75, "+126.8%", "البريد الترويجي", "وسيط 338 مقابل 149")
    add_metric(slide, 8.5, 3.55, 3.75, "+136.5%", "الصفحة الرئيسية", "وسيط 350 مقابل 148", accent=GOLD)
    add_card(slide, 8.5, 5.1, 3.75, 1.25, PALE)
    add_text(slide, "Mann–Whitney ثنائي الطرف\nتصحيح Holm لثلاث مقارنات", 8.75, 5.25, 3.2, 0.9, 12, NAVY, True)
    add_notes(
        slide,
        "بيّن أن المقارنة تقصر البيانات على أزواج المركز والوجبة التي لها سجلات مع الترويج وبدونه، لتخفيف اختلاف مزيج المنتجات. "
        "وسيط الطلب للبريد 338 مقابل 149 (فرق نسبي +126.8%)، وللصفحة 350 مقابل 148 (+136.5%). استُخدم Mann–Whitney ثنائي الطرف، وصُححت ثلاثة اختبارات بطريقة Holm. "
        "الاختيار غير عشوائي، وقد تتزامن الحملات مع الخصم أو مواسم وطلب سابق؛ لذلك هذا ارتباط لا أثر سببي ولا دليل ربحية. الخلية 10.",
    )

    slide = new_slide(
        deck,
        "08 | السؤال 5 — الخصم",
        "علاقة الخصم بالطلب موجبة لكنها ضعيفة",
        "ارتباط رتبي Spearman على سجلات التدريب؛ لا يقدّر أثر تغيير السعر.",
        9,
        "خلية الدفتر 11",
    )
    add_card(slide, 0.8, 1.9, 7.6, 4.75, WHITE, line=BLUE_PALE)
    add_figure(slide, "05_discount_relationship.png", 1.05, 2.15, 7.1, 4.15)
    add_metric(slide, 8.85, 2.2, 3.4, "0.122", "ارتباط Spearman", "موجب وضعيف")
    add_card(slide, 8.85, 3.85, 3.4, 2.3, PALE)
    add_text(slide, "كيف نفسر النتيجة؟", 9.1, 4.1, 2.9, 0.36, 15, NAVY, True)
    add_bullet(slide, "الترويج والخصم قد يتزامنان.", 9.05, 4.65, 2.9, h=0.48, size=11)
    add_bullet(slide, "اختيار الوجبة والأسبوع قد يختلف.", 9.05, 5.18, 2.9, h=0.48, size=11)
    add_bullet(slide, "لا تتوفر بيانات هامش الربح.", 9.05, 5.7, 2.9, h=0.42, size=11)
    add_notes(
        slide,
        "عرّف نسبة الخصم كما حُسبت: (base_price ناقص checkout_price) مقسومًا على base_price. ثم اشرح أن Spearman يقيس العلاقة الرتيبة، "
        "وأن القيمة 0.122 تمثل ارتباطًا موجبًا ضعيفًا. قيمة p الصغيرة في عينة كبيرة لا تجعل الأثر التجاري مهمًا. قد تترافق الخصومات مع العروض أو منتجات وأسْابيع مختلفة، "
        "ولا توجد تكاليف أو هوامش لتقدير الربحية. الخلية 11.",
    )

    slide = new_slide(
        deck,
        "09 | السؤال 6 — تذبذب الوجبات",
        "الأولوية للمراجعة تجمع بين الحجم والتذبذب",
        "الترتيب أداة فرز؛ وليس كمية شراء أو احتياطيًا جاهزًا.",
        10,
        "خلية الدفتر 12",
    )
    add_card(slide, 0.8, 1.9, 7.4, 4.75, WHITE, line=BLUE_PALE)
    add_figure(slide, "06_meal_volatility.png", 1.05, 2.15, 6.9, 4.15)
    add_metric(slide, 8.65, 2.0, 3.6, "2290", "أعلى أولوية مراجعة", "9.77M طلب • CV ≈ 0.90")
    add_card(slide, 8.65, 3.65, 3.6, 2.55, PALE)
    add_text(slide, "وجبات بارزة أخرى", 8.95, 3.9, 3.0, 0.36, 15, NAVY, True)
    add_text(slide, "1971   •   CV ≈ 1.10", 8.95, 4.55, 3.0, 0.38, 13, INK)
    add_text(slide, "1754   •   CV ≈ 0.89", 8.95, 5.1, 3.0, 0.38, 13, INK)
    add_text(slide, "راجع الطلب حسب المركز–الوجبة والمئين 90.", 8.95, 5.65, 3.0, 0.43, 11, TEAL, True)
    add_notes(
        slide,
        "اشرح أن الطلب جُمّع أسبوعيًا لكل وجبة عبر المراكز، ثم حُسب الانحراف المعياري مقسومًا على المتوسط (CV). "
        "رتّب الدفتر الوجبات باستخدام رتب مئينية للحجم والتذبذب معًا. ظهرت 2290 كأعلى أولوية، ثم 1971 و1754 ضمن الوجبات البارزة. "
        "لا تترجم هذا الترتيب مباشرة إلى مخزون: افحص السلسلة لكل مركز ووجبة، واستخدم المئين 90 كمرجع تجريبي، مع مراعاة الصلاحية والتغطية. الخلية 12.",
    )

    slide = new_slide(
        deck,
        "10 | إضافة اختيارية — التنبؤ",
        "خط الأساس البسيط تفوّق على الانحدار في التحقق",
        "مقارنة زمنية على الأسابيع 136–145، للصفوف ذات الأسبوع السابق المتصل.",
        11,
        "خلية الدفتر 13",
    )
    add_metric(slide, 0.95, 2.0, 3.5, "97.16", "MAE الانحدار الخطي", "طلب لكل سجل")
    add_metric(slide, 4.9, 2.0, 3.5, "81.78", "MAE خط الأسبوع السابق", "الأقل أفضل")
    add_metric(slide, 8.85, 2.0, 3.5, "32,573", "تنبؤًا محفوظًا", "صف لكل سجل اختبار")
    add_card(slide, 0.95, 3.75, 5.55, 2.3, PALE)
    add_text(slide, "ما الذي قورن؟", 1.25, 4.0, 4.8, 0.36, 16, NAVY, True)
    add_bullet(slide, "انحدار خطي بميزات الأسبوع والطلب السابق والسعر والترويج.", 1.2, 4.55, 4.9, h=0.55, size=12)
    add_bullet(slide, "خط أساس: تكرار طلب الأسبوع السابق.", 1.2, 5.25, 4.9, h=0.5, size=12)
    add_card(slide, 6.8, 3.75, 5.55, 2.3, WHITE, line=BLUE_PALE)
    add_text(slide, "كيف نستخدم النتيجة؟", 7.1, 4.0, 4.8, 0.36, 16, NAVY, True)
    add_bullet(slide, "اختير خط الأساس لأنه حقق MAE أقل في التحقق.", 7.05, 4.55, 4.9, h=0.55, size=12)
    add_bullet(slide, "التنبؤات محفوظة في reports/test_predictions.csv.", 7.05, 5.25, 4.9, h=0.5, size=12)
    add_text(slide, "تنبيه: التحقق خطوة واحدة للأمام، وليس محاكاة مستقلة لعشرة أسابيع متتابعة.", 1.1, 6.42, 11.2, 0.35, 11, GOLD, True)
    add_notes(
        slide,
        "هذه إضافة اختيارية في الدفتر. قُسم التدريب زمنيًا، واستخدمت آخر عشرة أسابيع للتحقق. انحدار NumPy الخطي استخدم الطلب السابق، الأسبوع، الخصم، ومؤشري الترويج؛ "
        "وقورن بخط أساس يكرر طلب الأسبوع السابق على السجلات ذات الأسبوع السابق المتصل. MAE هو متوسط الخطأ المطلق لكل سجل: 97.16 للانحدار و81.78 للخط الأساس، لذا اختير الخط الأساس. "
        "أُعيد التدريب على بيانات التدريب لإنتاج تنبؤ لكل صف اختبار، ثم تُرحّل التنبؤات أسبوعًا بعد أسبوع. لا نملك قيم الطلب الفعلية للاختبار لتقييمها، والتحقق ليس محاكاة عشرة أسابيع كاملة. الخلية 13.",
    )

    slide = new_slide(
        deck,
        "11 | من النتيجة إلى الإجراء",
        "ثلاث تجارب عملية — وثلاثة مقاييس نجاح",
        "ابدأ بتجارب صغيرة قابلة للمقارنة، ولا تعمم النتائج الرصدية مباشرة.",
        12,
        "خلية الدفتر 14",
    )
    recommendations = [
        (
            "01",
            "تجربة مخزون حسب المركز–الوجبة",
            "ابدأ بالوجبات مرتفعة الحجم والتذبذب؛ استخدم المئين 90 مرجعًا تجريبيًا.",
            "النجاح: خفض النفاد والهدر لكل 1,000 طلب دون تراجع الخدمة.",
        ),
        (
            "02",
            "اختبار عشوائي للترويج",
            "افصل البريد عن الصفحة، وحدد مجموعة مقارنة وفترة ومقياسًا مسبقًا.",
            "النجاح: طلب وهامش إضافيان بعد التكلفة؛ مع مراقبة النفاد والإلغاءات.",
        ),
        (
            "03",
            "اختبار درجات الخصم",
            "اختبر السعر على وجبات/مراكز محددة مع توثيق العروض المتزامنة.",
            "النجاح: هامش مساهمة إضافي بعد تكلفة المواد والحملة والهدر.",
        ),
    ]
    for i, (number, title, action, metric) in enumerate(recommendations):
        x = 0.85 + i * 4.0
        add_card(slide, x, 2.15, 3.72, 3.95, PALE if i != 1 else BLUE_PALE)
        add_text(slide, number, x + 0.22, 2.42, 0.65, 0.5, 23, TEAL, True, rtl=False)
        add_text(slide, title, x + 0.22, 3.05, 3.15, 0.75, 16, NAVY, True)
        add_text(slide, action, x + 0.22, 3.95, 3.15, 0.9, 12, INK)
        add_text(slide, metric, x + 0.22, 5.0, 3.15, 0.86, 11, TEAL, True)
    add_notes(
        slide,
        "اعرض التوصيات باعتبارها خطوات لا نتائج مثبتة. تجربة المخزون تستفيد من التذبذب والطلب حسب المركز والوجبة، وتُقاس بمعدل النفاد والهدر لكل ألف طلب مع مستوى الخدمة. "
        "اختبار الترويج يجب أن يحدد وحدة التوزيع والعلاج والمقارنة والفترة مسبقًا، ويفصل القنوات قدر الإمكان. "
        "اختبار الخصم يحتاج حساب هامش المساهمة وتكاليف المواد والحملة. الخلية 14 وتوصيات تقرير العميل.",
    )

    slide = new_slide(
        deck,
        "12 | الحدود والخلاصة",
        "ما الذي تقوله البيانات — وما الذي لا تستطيع إثباته؟",
        "النتيجة الأقوى: وصف الطلب المسجل وتحديد ما يستحق اختبارًا تشغيليًا.",
        13,
        "خلايا الدفتر 1 و14–15",
    )
    add_card(slide, 0.85, 2.0, 5.6, 3.75, PALE)
    add_text(slide, "حدود يجب تذكرها", 1.15, 2.28, 4.8, 0.4, 17, NAVY, True)
    limits = [
        "الأسابيع أرقام وليست تواريخ؛ والمواقع مجهولة الهوية.",
        "الطلب المسجل قد يكون مقيدًا بتوفر المخزون.",
        "لا بيانات مباشرة عن النفاد أو الهدر أو التكاليف.",
        "العروض والأسعار غير عشوائية؛ الارتباط لا يثبت السببية.",
        "غياب سجل لا يعني أن الطلب يساوي صفرًا.",
    ]
    for i, item in enumerate(limits):
        add_bullet(slide, item, 1.1, 2.9 + i * 0.5, 4.95, h=0.42, size=11, marker=GOLD)
    add_card(slide, 6.8, 2.0, 5.65, 3.75, WHITE, line=BLUE_PALE)
    add_text(slide, "منهج العمل والتحقق", 7.1, 2.28, 4.85, 0.4, 17, NAVY, True)
    add_bullet(slide, "Copilot ساعد في إعداد سير التحليل والتوثيق.", 7.05, 2.95, 4.95, h=0.48, size=12)
    add_bullet(slide, "الأرقام والتوصيات مرتبطة بخلايا الدفتر ونتائجه.", 7.05, 3.6, 4.95, h=0.48, size=12)
    add_bullet(slide, "أُعيد تشغيل الخلايا والتحقق من الرسوم الستة والمخرجات.", 7.05, 4.25, 4.95, h=0.55, size=12)
    add_bullet(slide, "التنبؤات تجريبية وتحتاج تحققًا بعد ظهور القيم الفعلية.", 7.05, 4.95, 4.95, h=0.5, size=12)
    add_text(slide, "الخلاصة: استخدم التحليل لتحديد أين نختبر، لا لافتراض ما الذي سبّب النتيجة.", 1.05, 6.15, 11.2, 0.48, 15, TEAL, True)
    add_notes(
        slide,
        "اختم بالفصل بين الوصف والاستنتاج: نعرف كيف توزّع الطلب المسجل في هذه البيانات، لكننا لا نعرف مباشرةً الطلب غير الملبّى أو الهدر أو الربحية. "
        "لا توجد تواريخ فعلية أو مواقع معلنة، والترويج ليس تجربة عشوائية. اذكر أن Copilot استُخدم للمساعدة في إعداد التحليل والتوثيق، "
        "لكن النتائج أعيد إنتاجها من الدفتر وتم التحقق من ملفات الرسوم والمخرجات. ادعُ إلى اختبار مضبوط قبل تغيير سياسة المخزون أو التسعير. الخلايا 14–15.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    print(
        f"Saved {len(deck.slides)} slides to {OUTPUT.relative_to(ROOT)} "
        f"({OUTPUT.stat().st_size:,} bytes)"
    )


if __name__ == "__main__":
    main()
